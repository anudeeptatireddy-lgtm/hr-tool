"""Text extraction by file format. Never raises: a bad file returns ok=False with an error."""
import logging
import subprocess
from dataclasses import dataclass
from pathlib import Path

log = logging.getLogger(__name__)

SUPPORTED = {".pdf", ".docx", ".doc", ".txt", ".md", ".rtf"}
# Below this many characters a CV is probably a scanned image or an empty file.
MIN_CHARS = 200


@dataclass
class Extracted:
    path: Path
    fmt: str
    text: str
    ok: bool
    error: str = ""

    @property
    def chars(self) -> int:
        return len(self.text)


def _pdf(path: Path) -> str:
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        text = "\n".join((page.extract_text() or "") for page in pdf.pages)
    if text.strip():
        return text
    # Fallback: pypdf sometimes reads PDFs pdfplumber can't.
    from pypdf import PdfReader

    return "\n".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)


def _docx(path: Path) -> str:
    import docx

    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    # CVs often put roles or skills in tables; python-docx skips those in .paragraphs.
    for table in doc.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                t = cell.text.strip()
                if t and t not in cells:  # merged cells repeat their text
                    cells.append(t)
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _via_textutil(path: Path) -> str:
    # macOS built-in; converts legacy .doc and .rtf to plain text.
    out = subprocess.run(
        ["textutil", "-convert", "txt", "-stdout", str(path)],
        capture_output=True, text=True, timeout=30,
    )
    if out.returncode != 0:
        raise RuntimeError(out.stderr.strip() or "textutil failed")
    return out.stdout


def _txt(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8", "utf-16", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    lines = [" ".join(line.split()) for line in text.split("\n")]
    out, blank = [], 0
    for line in lines:
        blank = blank + 1 if not line else 0
        if blank <= 1:
            out.append(line)
    return "\n".join(out).strip()


def extract(path: Path) -> Extracted:
    path = Path(path)
    fmt = path.suffix.lower().lstrip(".")
    try:
        if path.suffix.lower() == ".pdf":
            text = _pdf(path)
        elif path.suffix.lower() == ".docx":
            text = _docx(path)
        elif path.suffix.lower() in {".doc", ".rtf"}:
            text = _via_textutil(path)
        elif path.suffix.lower() in {".txt", ".md"}:
            text = _txt(path)
        else:
            return Extracted(path, fmt, "", False, f"unsupported format .{fmt}")
        text = _normalise(text)
        if len(text) < MIN_CHARS:
            return Extracted(path, fmt, text, False, f"too little text ({len(text)} chars); scanned or empty?")
        return Extracted(path, fmt, text, True)
    except Exception as exc:  # never crash the batch
        log.warning("extract failed for %s: %s", path.name, exc)
        return Extracted(path, fmt, "", False, f"{type(exc).__name__}: {exc}")
